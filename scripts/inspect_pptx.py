from pptx import Presentation
import os

pptx_path = os.path.join("presentation", "FinSight_AI_Pitch_Deck_Enterprise.pptx")
prs = Presentation(pptx_path)

out_lines = []
out_lines.append(f"Total slides: {len(prs.slides)}\n")

for idx, slide in enumerate(prs.slides):
    out_lines.append(f"\n==================== SLIDE {idx+1} ====================")
    for shape_idx, shape in enumerate(slide.shapes):
        shape_type = str(shape.shape_type)
        pos = f"Left={shape.left.inches:.2f}, Top={shape.top.inches:.2f}, W={shape.width.inches:.2f}, H={shape.height.inches:.2f}"
        if shape.has_text_frame:
            full_text = "\n".join([p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip()])
            if full_text:
                out_lines.append(f"  Shape {shape_idx} ({pos}) [Text]:\n{full_text}")
        elif shape.has_table:
            table_lines = []
            for row in shape.table.rows:
                cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                table_lines.append(" | ".join(cells))
            out_lines.append(f"  Shape {shape_idx} ({pos}) [Table]:\n" + "\n".join(table_lines))
        elif shape.shape_type == 13: # Picture
            out_lines.append(f"  Shape {shape_idx} ({pos}) [Picture]: {shape.name}")
        else:
            if hasattr(shape, "text") and shape.text:
                out_lines.append(f"  Shape {shape_idx} ({pos}) [Other]: {shape.text}")

out_path = os.path.join("presentation", "enterprise_deck_dump.txt")
with open(out_path, "w", encoding="utf-8") as f:
    f.write("\n".join(out_lines))

print(f"Dump written to {out_path}")
