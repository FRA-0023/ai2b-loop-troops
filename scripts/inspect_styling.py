from pptx import Presentation
import os

prs = Presentation(os.path.join("presentation", "FinSight_AI_Pitch_Deck_Enterprise.pptx"))
for idx, slide in enumerate(prs.slides):
    print(f"--- Slide {idx+1} ---")
    for s_idx, shape in enumerate(slide.shapes):
        if s_idx >= 6:
            break
        txt = ""
        font_info = ""
        if shape.has_text_frame and shape.text_frame.text:
            txt = shape.text_frame.text.replace("\n", " ")[:30]
            p = shape.text_frame.paragraphs[0]
            if p.runs:
                r = p.runs[0]
                sz = r.font.size.pt if r.font.size else None
                font_info = f"font={r.font.name}, sz={sz}"
        print(f"  Shape {s_idx}: {shape.name} {font_info} -> '{txt}'")
